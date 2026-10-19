"""Check the Craft live E2E outputs outside the servers, then render each with its own editor.

python e2e_check.py <out dir>      (the folder web/e2e/craft-live.spec.ts filled)

The checks run in aura-sandbox:craft-spike, the box image from before the change, which still has
LibreOffice and the Python Office libraries: nothing a Craft server wrote is read back by a Craft
server. The renders are the editors' own headless views of the result (the "screenshot" each
editor gives the agent), taken in the identity's box under the recipes' landrun fence.
"""
import base64
import glob
import json
import os
import subprocess
import sys

OUT = os.path.abspath(sys.argv[1])
os.environ.setdefault("CRAFT_WRAP", "landrun --best-effort --rox / --rw /workspace --rw /dev "
                      "--unrestricted-network --unrestricted-scoped --env HOME --env PATH --")
from mcpbox import BOX, Session  # noqa: E402

BOXDIR = "/workspace/craft-e2e"
WORD_PAGE = int(os.environ.get("CRAFT_WORD_PAGE", "2"))

CHECKS = r"""
import glob, json, os, subprocess, zipfile
from statistics import mean
out = {}
def one(pattern):
    found = sorted(glob.glob('/out/' + pattern))
    return found[-1] if found else None

f = one('gridcraft-*.xlsx')
if f:
    import openpyxl
    v = openpyxl.load_workbook(f, data_only=True)
    os.makedirs('/tmp/re', exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', '/tmp/re', f],
                   capture_output=True)
    r = openpyxl.load_workbook('/tmp/re/' + os.path.basename(f), data_only=True)
    out['gridcraft'] = {'file': os.path.basename(f), 'risk_free_B39': v['Master Inputs Start here']['B39'].value,
                        'value_per_share_cached_F71': v['Valuation Model']['F71'].value,
                        'value_per_share_libreoffice_F71': r['Valuation Model']['F71'].value,
                        'expected_F71_at_5pct': 333.619261987673}

f = one('pdfcraft-*.pdf')
if f:
    text = subprocess.run(['pdftotext', f, '-'], capture_output=True, text=True).stdout
    orig = subprocess.run(['pdftotext', '/real/anac-nota-illustrativa.pdf', '-'], capture_output=True, text=True).stdout
    pages = [l for l in subprocess.run(['pdfinfo', f], capture_output=True, text=True).stdout.splitlines() if l.startswith('Pages')]
    out['pdfcraft'] = {'file': os.path.basename(f), 'ANAC_in_original': orig.count('ANAC'), 'ANAC_left': text.count('ANAC'),
                       'words_original': len(orig.split()), 'words': len(text.split()), 'pages': pages}

f = one('wordcraft-*.docx')
if f:
    import docx
    d, o = docx.Document(f), docx.Document('/real/bando-tipo-anac-2.docx')
    t = '\n'.join(p.text for p in d.paragraphs) + '\n'.join(c.text for tb in d.tables for row in tb.rows for c in row.cells)
    out['wordcraft'] = {'file': os.path.basename(f), 'has_125_2026': '125/2026' in t, 'has_12_ottobre_2026': '12 ottobre 2026' in t,
                        'placeholder_left': 'a contrarre n. ....' in t,
                        'paragraphs': len(d.paragraphs), 'paragraphs_original': len(o.paragraphs),
                        'tables': len(d.tables), 'tables_original': len(o.tables),
                        'bold_runs': sum(1 for p in d.paragraphs for r in p.runs if r.bold),
                        'bold_runs_original': sum(1 for p in o.paragraphs for r in p.runs if r.bold)}

f = one('deckcraft-*.pptx')
if f:
    from pptx import Presentation
    p, o = Presentation(f), Presentation('/real/HWC20_PPT_IT.pptx')
    texts = [[s.text_frame.text for s in sl.shapes if s.has_text_frame] for sl in p.slides]
    out['deckcraft'] = {'file': os.path.basename(f), 'slides': len(p.slides), 'slides_original': len(o.slides),
                        'last_slide_text': texts[-1], 'domande_on_slide': [k + 1 for k, t in enumerate(texts) if any('Domande?' in x for x in t)]}

f = one('photocraft-*.jpg')
if f:
    from PIL import Image, ImageStat
    i, o = Image.open(f), Image.open('/real/ponte-vecchio-dusk.jpg')
    o2 = o.resize(i.size)
    out['photocraft'] = {'file': os.path.basename(f), 'size': i.size, 'size_original': o.size,
                         'exif': bool(i.info.get('exif')), 'exif_original': bool(o.info.get('exif')),
                         'mean_luma': round(ImageStat.Stat(i.convert('L')).mean[0], 1),
                         'mean_luma_original_resized': round(ImageStat.Stat(o2.convert('L')).mean[0], 1)}

f = one('filmcraft-*.mp4')
if f:
    data = open(f, 'rb').read()
    i = data.find(b'mvhd'); ver = data[i + 4]
    ts, dur = (int.from_bytes(data[i+16:i+20], 'big'), int.from_bytes(data[i+20:i+24], 'big')) if ver == 0 else \
              (int.from_bytes(data[i+24:i+28], 'big'), int.from_bytes(data[i+28:i+36], 'big'))
    sizes, tk = [], data.find(b'tkhd')
    while tk != -1:  # tkhd: width and height, 16.16 fixed, end the box (ISO/IEC 14496-12 8.3.2)
        off = tk + (80 if data[tk + 4] == 0 else 92)
        sizes.append([int.from_bytes(data[off:off+4], 'big') >> 16, int.from_bytes(data[off+4:off+8], 'big') >> 16])
        tk = data.find(b'tkhd', tk + 4)
    out['filmcraft'] = {'file': os.path.basename(f), 'bytes': len(data), 'seconds': round(dur / ts, 2),
                        'avc1': b'avc1' in data, 'track_sizes': sizes}
print(json.dumps(out, indent=1, default=str))
"""


def run_checks(real):
    proc = subprocess.run(["docker", "run", "--rm", "-i", "--network", "none", "-v", f"{OUT}:/out:ro", "-v", f"{real}:/real:ro",
                           "--entrypoint", "python3", "aura-sandbox:craft-spike", "-"],
                          input=CHECKS.encode("utf-8"), capture_output=True, env={**os.environ, "MSYS_NO_PATHCONV": "1"})
    text = proc.stdout.decode("utf-8", "replace")
    if proc.returncode != 0:
        raise SystemExit(proc.stderr.decode("utf-8", "replace")[-3000:])
    with open(os.path.join(OUT, "checks.json"), "w", encoding="utf-8") as fh:
        fh.write(text)
    print(text)


def shot(session, tool, **args):
    """The base64 PNGs a render tool answered (Session.call keeps only their lengths, for the log)."""
    reply, _ = session.request("tools/call", {"name": tool, "arguments": args})
    images = [c["data"] for c in (reply.get("result") or {}).get("content", []) if c.get("type") == "image"]
    if not images:
        raise RuntimeError(f"{session.server} {tool}: no image: {json.dumps(reply)[:400]}")
    return images


def save_png(name, images_b64):
    path = os.path.join(OUT, f"{name}-render.png")
    with open(path, "wb") as fh:
        fh.write(base64.b64decode(images_b64[0]))
    print("render:", path)


def render():
    subprocess.run(["docker", "exec", BOX, "mkdir", "-p", BOXDIR], check=True)
    files = {}
    for path in glob.glob(os.path.join(OUT, "*-*.*")):
        server, _, name = os.path.basename(path).partition("-")
        if server in ("gridcraft", "pdfcraft", "wordcraft", "deckcraft", "photocraft", "filmcraft") and not name.endswith((".png", ".json", ".log")):
            subprocess.run(["docker", "cp", path, f"{BOX}:{BOXDIR}/{os.path.basename(path)}"], check=True)
            files[server] = f"{BOXDIR}/{os.path.basename(path)}"

    if "pdfcraft" in files or "gridcraft" in files:
        pdf = Session("pdfcraft")
        if "gridcraft" in files:
            g = Session("gridcraft")
            g.call("open_workbook", path=files["gridcraft"])
            sheet = f"{BOXDIR}/gridcraft-valuation.pdf"
            g.call("execute_command", command="sheet.activate", params={"sheet": "Valuation Model"})
            g.call("execute_command", command="file.exportPdf",
                   params={"path": sheet, "range": "A55:F75", "fitToPage": True})
            g.close()
            doc = pdf.json("doc_open", path=sheet)
            save_png("gridcraft", shot(pdf, "render_preview", doc=doc["doc"] if isinstance(doc, dict) else 1, page=1, max_side=1400))
        if "pdfcraft" in files:
            doc = pdf.json("doc_open", path=files["pdfcraft"])
            save_png("pdfcraft", shot(pdf, "render_preview", doc=doc["doc"] if isinstance(doc, dict) else 2, page=3, max_side=1400))
        pdf.close()
    if "wordcraft" in files:
        w = Session("wordcraft")
        w.call("open_document", path=files["wordcraft"])
        w.call("select_text", text="125/2026")
        page = WORD_PAGE
        save_png("wordcraft", shot(w, "render_page", page=page, scale=1.5))
        w.close()
    if "deckcraft" in files:
        d = Session("deckcraft")
        d.call("open_presentation", path=files["deckcraft"])
        slides = (d.json("inspect_document") or {}).get("slides", [])
        save_png("deckcraft", shot(d, "render_slide", slide=max(len(slides) - 1, 0), scale=1.5))
        d.close()
    if "photocraft" in files:
        p = Session("photocraft")
        p.call("doc_open", path=files["photocraft"])
        save_png("photocraft", shot(p, "render_preview", max_side=1400))
        p.close()
    if "filmcraft" in files:
        f = Session("filmcraft")
        item = f.json("media_import", paths=[files["filmcraft"]])["items"][0]
        f.call("command_run", id="file.newSequence", params={"name": "Check", "fromItem": item})
        save_png("filmcraft", shot(f, "render_frame", seconds=5, max_side=1280))
        f.close()


if __name__ == "__main__":
    run_checks(os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else r"D:\tmp\craft-real")
    render()
