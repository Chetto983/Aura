"""One real task per Craft server, an independent check of what it wrote, and today's path.

python -I tasks.py <server>   -> runs/<server>.json
"""
import json
import os
import sys
import time

from mcpbox import Session, box

REAL = "/workspace/craft-spike/real"
OUT = "/workspace/craft-spike/out"


def gridcraft():
    """Damodaran's FCFF model (16 sheets, 906 formulas): raise the riskfree rate from 4.5% to 5% and
    read the value per share; the saved xlsx is recalculated by LibreOffice as the oracle."""
    s = Session("gridcraft")
    r = {"file": "fcffginzu.xls"}
    t = time.monotonic()
    s.call("open_workbook", path=f"{REAL}/fcffginzu.xls")
    r["inspect"] = s.json("inspect_workbook")
    r["before"] = s.json("read_range", sheet="Valuation Model", range="F71")
    s.call("set_cell", cell="'Master Inputs Start here'!B39", input="0.05")
    r["riskfree_after"] = s.json("read_range", sheet="Master Inputs Start here", range="B39")
    r["after"] = s.json("read_range", sheet="Valuation Model", range="F71")
    s.call("save_workbook", path=f"{OUT}/fcffginzu-rf5.xlsx")
    r["craft_seconds"] = round(time.monotonic() - t, 2)
    r["peak_kb"] = s.close()
    r["calls"] = s.log
    r["check"] = box(f"""cd /tmp && rm -rf chk && mkdir chk && cd chk
python3 - <<'PY'
import openpyxl
wb=openpyxl.load_workbook('{OUT}/fcffginzu-rf5.xlsx'); wv=openpyxl.load_workbook('{OUT}/fcffginzu-rf5.xlsx', data_only=True)
f=sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('='))
print('saved: sheets', len(wb.worksheets), 'formulas', f, 'B39', wv['Master Inputs Start here']['B39'].value, 'cached F71', wv['Valuation Model']['F71'].value)
PY
soffice --headless --convert-to xlsx:"Calc MS Excel 2007 XML" --outdir /tmp/chk {OUT}/fcffginzu-rf5.xlsx >/dev/null 2>&1
python3 -c "import openpyxl; w=openpyxl.load_workbook('/tmp/chk/fcffginzu-rf5.xlsx', data_only=True); print('LibreOffice recalculated F71', w['Valuation Model']['F71'].value)"
""")
    r["baseline"] = box(f"""cd /tmp && rm -rf base && mkdir base && cd base
s=$(date +%s%N)
soffice --headless --convert-to xlsx --outdir /tmp/base {REAL}/fcffginzu.xls >/dev/null 2>&1
python3 - <<'PY'
import openpyxl
wb=openpyxl.load_workbook('/tmp/base/fcffginzu.xlsx'); wb['Master Inputs Start here']['B39']=0.05; wb.save('/tmp/base/edited.xlsx')
print('openpyxl cached F71 after its own save', openpyxl.load_workbook('/tmp/base/edited.xlsx', data_only=True)['Valuation Model']['F71'].value)
PY
mkdir -p /tmp/base/re && soffice --headless --convert-to xlsx --outdir /tmp/base/re /tmp/base/edited.xlsx >/dev/null 2>&1
echo "baseline ms $(( ($(date +%s%N)-s)/1000000 ))"
python3 -c "import openpyxl; print('baseline LibreOffice F71', openpyxl.load_workbook('/tmp/base/re/edited.xlsx', data_only=True)['Valuation Model']['F71'].value)"
""")
    return r


def pdfcraft():
    """ANAC's 42-page note: redact every "ANAC" for good, audit it, save, export to Word. Then OCR three
    pages of a 352-page public-domain scan with no text layer, against tesseract on the same pages."""
    s = Session("pdfcraft")
    r = {"files": ["anac-nota-illustrativa.pdf", "inventari-manoscritti-37.pdf"]}
    t = time.monotonic()
    doc = s.json("doc_open", path=f"{REAL}/anac-nota-illustrativa.pdf")
    d = doc["doc"] if isinstance(doc, dict) else 1
    text = s.json("text_extract", doc=d)
    r["extracted_chars"] = len(json.dumps(text))
    r["found"] = s.json("text_find", doc=d, query="ANAC", limit=50)
    r["mark"] = s.json("tool_call", name="redact_mark", arguments={"doc": d, "find": "ANAC"})
    r["apply"] = s.json("tool_call", name="redact_apply", arguments={"doc": d})
    r["audit"] = s.json("tool_call", name="doc_audit_redactions", arguments={"doc": d})
    s.call("doc_save", doc=d, path=f"{OUT}/anac-redatto.pdf")
    r["export"] = s.json("tool_call", name="doc_export_office", arguments={"doc": d, "path": f"{OUT}/anac-nota.docx"})
    r["redact_export_seconds"] = round(time.monotonic() - t, 2)
    t = time.monotonic()
    scan = s.json("doc_open", path=f"{REAL}/inventari-manoscritti-37.pdf")
    sd = scan["doc"] if isinstance(scan, dict) else 2
    r["ocr"] = s.json("tool_call", name="ocr_recognize", arguments={"doc": sd, "pages": [10, 11, 12]})
    r["ocr_seconds"] = round(time.monotonic() - t, 2)
    ocr_text = s.json("text_extract", doc=sd, pages=[10, 11, 12])
    r["ocr_text_sample"] = json.dumps(ocr_text, ensure_ascii=False)[:1200]
    s.call("doc_save", doc=sd, path=f"{OUT}/inventari-ocr.pdf")
    r["peak_kb"] = s.close()
    r["calls"] = s.log
    r["check"] = box(f"""cd /tmp
echo "ANAC in the original: $(pdftotext {REAL}/anac-nota-illustrativa.pdf - | grep -o ANAC | wc -l), in the redacted: $(pdftotext {OUT}/anac-redatto.pdf - | grep -o ANAC | wc -l)"
echo "pages $(pdfinfo {OUT}/anac-redatto.pdf | grep Pages)"
echo "words in the original $(pdftotext {REAL}/anac-nota-illustrativa.pdf - | wc -w), in the redacted $(pdftotext {OUT}/anac-redatto.pdf - | wc -w)"
python3 -c "import docx; d=docx.Document('{OUT}/anac-nota.docx'); t=chr(10).join(p.text for p in d.paragraphs); print('export docx: paragraphs', len(d.paragraphs), 'tables', len(d.tables), 'words', len(t.split()), 'has Sommario', 'Sommario' in t)"
echo "searchable text on scan pages 10-12 after OCR: $(pdftotext -f 10 -l 12 {OUT}/inventari-ocr.pdf - | wc -w) words"
""")
    r["baseline"] = box(f"""cd /tmp && rm -rf ocrb && mkdir ocrb && cd ocrb
s=$(date +%s%N)
pdftoppm -f 10 -l 12 -r 300 -png {REAL}/inventari-manoscritti-37.pdf p
for f in p-*.png; do tesseract $f ${{f%.png}} -l eng >/dev/null 2>&1; done
echo "tesseract eng, 3 pages at 300 dpi: $(( ($(date +%s%N)-s)/1000000 )) ms, $(cat p-*.txt | wc -w) words"
echo "sample: $(cat p-*.txt | tr '\n' ' ' | cut -c1-600)"
s=$(date +%s%N); soffice --headless --infilter=writer_pdf_import --convert-to docx --outdir /tmp/ocrb {REAL}/anac-nota-illustrativa.pdf >/dev/null 2>&1
echo "LibreOffice pdf->docx: $(( ($(date +%s%N)-s)/1000000 )) ms"
python3 -c "import docx; d=docx.Document('/tmp/ocrb/anac-nota-illustrativa.docx'); t=chr(10).join(p.text for p in d.paragraphs); print('LO docx: paragraphs', len(d.paragraphs), 'tables', len(d.tables), 'words', len(t.split()))"
echo "true redaction tool in the box: none (pypdf, qpdf and pdftk draw or remove nothing under a mark)"
""")
    return r


from tasks_more import deckcraft, filmcraft, photocraft, wordcraft  # noqa: E402

TASKS = {"gridcraft": gridcraft, "pdfcraft": pdfcraft, "wordcraft": wordcraft,
         "deckcraft": deckcraft, "filmcraft": filmcraft, "photocraft": photocraft}

if __name__ == "__main__":
    name = sys.argv[1]
    result = TASKS[name]()
    json.dump(result, open(f"{os.environ.get('CRAFT_RUNS', 'runs')}/{name}.json", "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    for c in result.get("calls", []):
        print(f"  {c['tool']:<22} {'ERR' if c['error'] else 'ok':<4} {c['ms']:>6} ms  {c['text'][:150]!r}")
    for k, v in result.items():
        if k not in ("calls", "inspect", "project"):
            print(f"{k}: {str(v)[:600]}")
