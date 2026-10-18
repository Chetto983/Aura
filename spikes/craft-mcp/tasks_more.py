"""WordCraft, DeckCraft, FilmCraft and PhotoCraft tasks (see tasks.py)."""
import json
import re
import time

from mcpbox import Session, box

REAL = "/workspace/craft-spike/real"
OUT = "/workspace/craft-spike/out"
FILL = "a contrarre n. 125/2026 del 12 ottobre 2026"
TITLE = "Ambienti di lavoro sani e sicuri - Comune di Trento 2026"


def wordcraft():
    """ANAC's bando tipo n. 2 (821 paragraphs, 12 tables): fill the decree's number and date, a fill-in
    spread over 19 runs, then save as docx and pdf and look at page 1."""
    s = Session("wordcraft")
    r = {"file": "bando-tipo-anac-2.docx"}
    t = time.monotonic()
    s.call("open_document", path=f"{REAL}/bando-tipo-anac-2.docx")
    text = s.json("get_text")
    body = text if isinstance(text, str) else json.dumps(text, ensure_ascii=False)
    m = re.search(r"a contrarre n\. [.]+ del [….]+", body)
    r["target"] = m.group(0) if m else None
    if m:
        s.call("select_text", text=m.group(0))
        s.call("type_text", text=FILL)
    s.call("save_document", path=f"{OUT}/bando-compilato.docx")
    s.call("save_document", path=f"{OUT}/bando-compilato.pdf")
    _, images, _ = s.call("render_page", page=1)
    r["render_page_png_b64"] = images
    r["craft_seconds"] = round(time.monotonic() - t, 2)
    r["peak_kb"] = s.close()
    r["calls"] = s.log
    r["check"] = box(f"""python3 - <<'PY'
import docx
for n in ['{REAL}/bando-tipo-anac-2.docx', '{OUT}/bando-compilato.docx']:
    d=docx.Document(n); t=chr(10).join(p.text for p in d.paragraphs)
    print(n.rsplit('/',1)[1], 'paragraphs', len(d.paragraphs), 'tables', len(d.tables), 'filled', '{FILL}' in t,
          'bold runs', sum(1 for p in d.paragraphs for x in p.runs if x.bold))
PY
echo "craft pdf: $(pdfinfo {OUT}/bando-compilato.pdf | grep Pages), filled in its text: $(pdftotext {OUT}/bando-compilato.pdf - | tr '\\n' ' ' | grep -c '125/2026')"
""")
    r["baseline"] = box(f"""cd /tmp && rm -rf wb && mkdir wb && cd wb
s=$(date +%s%N)
python3 - <<'PY'
import docx, re
d=docx.Document('{REAL}/bando-tipo-anac-2.docx')
p=next(p for p in d.paragraphs if 'a contrarre n.' in p.text)
print('fill-in inside one run:', any(re.search(r'a contrarre n\\. [.]+ del', x.text) for x in p.runs), '- runs:', len(p.runs))
styles=sorted({{(bool(x.bold), bool(x.italic), bool(x.underline)) for x in p.runs}})
p.runs[0].text=re.sub(r'a contrarre n\\. [.]+ del [….]+', '{FILL}', p.text)
for x in p.runs[1:]: x.text=''
print('run styles before', styles, '- after a paragraph-level replace every word takes the first run style')
d.save('/tmp/wb/bando-python.docx')
PY
soffice --headless --convert-to pdf --outdir /tmp/wb /tmp/wb/bando-python.docx >/dev/null 2>&1
echo "python-docx + LibreOffice: $(( ($(date +%s%N)-s)/1000000 )) ms, $(pdfinfo /tmp/wb/bando-python.pdf | grep Pages)"
""")
    return r


def deckcraft():
    """EU-OSHA's Italian deck (15 slides, 23 pictures): retitle the first slide, look at it, export to
    pdf and save the pptx."""
    s = Session("deckcraft")
    r = {"file": "HWC20_PPT_IT.pptx"}
    t = time.monotonic()
    s.call("open_presentation", path=f"{REAL}/HWC20_PPT_IT.pptx")
    r["inspect"] = s.json("inspect_document")
    slide = s.json("inspect_slide", index=0)
    shapes = slide.get("shapes", []) if isinstance(slide, dict) else []
    texts = [x for x in shapes if (x.get("text") or "").strip()]
    r["first_text_shape"] = texts[0] if texts else slide
    if texts:
        s.call("go_to_slide", index=0)
        s.call("set_text", id=texts[0]["id"], text=TITLE)
    _, images, _ = s.call("render_slide", slide=0)
    r["render_slide_png_b64"] = images
    s.call("export", path=f"{OUT}/osha.pdf", format="pdf")
    s.call("save_presentation", path=f"{OUT}/osha-edit.pptx")
    r["craft_seconds"] = round(time.monotonic() - t, 2)
    r["peak_kb"] = s.close()
    r["calls"] = s.log
    r["check"] = box(f"""python3 - <<'PY'
import pptx
for n in ['{REAL}/HWC20_PPT_IT.pptx', '{OUT}/osha-edit.pptx']:
    p=pptx.Presentation(n)
    pics=sum(1 for s in p.slides for sh in s.shapes if sh.shape_type==13)
    first=[sh.text_frame.text for sh in p.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    print(n.rsplit('/',1)[1], 'slides', len(p.slides), 'pictures', pics, 'first slide text', first[:2])
PY
echo "craft pdf: $(pdfinfo {OUT}/osha.pdf | grep Pages), retitled in its text: $(pdftotext -l 1 {OUT}/osha.pdf - | grep -c Trento)"
""")
    r["baseline"] = box(f"""cd /tmp && rm -rf db && mkdir db && cd db
s=$(date +%s%N)
python3 - <<'PY'
import pptx
p=pptx.Presentation('{REAL}/HWC20_PPT_IT.pptx')
sh=next(sh for sh in p.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip())
runs=[r for para in sh.text_frame.paragraphs for r in para.runs]
runs[0].text='{TITLE}'
for x in runs[1:]: x.text=''
p.save('/tmp/db/osha-python.pptx')
PY
soffice --headless --convert-to pdf --outdir /tmp/db /tmp/db/osha-python.pptx >/dev/null 2>&1
echo "python-pptx + LibreOffice: $(( ($(date +%s%N)-s)/1000000 )) ms, $(pdfinfo /tmp/db/osha-python.pdf | grep Pages)"
""")
    return r


def filmcraft():
    """Blender's Sintel trailer (52 s, 1080p): a sequence from the clip, one frame looked at, and an
    H.264 1280x544 export, headless, at the box's limits. The box has no ffmpeg."""
    s = Session("filmcraft")
    r = {"file": "sintel-trailer-1080p.mp4"}
    t = time.monotonic()
    r["import"] = s.json("media_import", paths=[f"{REAL}/sintel-trailer-1080p.mp4"])
    project = s.json("project_inspect")
    r["project"] = project
    item = (r["import"] or {}).get("items", [None])[0]
    r["item"] = item
    r["sequence"] = s.json("command_run", id="file.newSequence", params={"name": "Trailer", "fromItem": item})
    _, images, _ = s.call("render_frame", seconds=20, max_side=640)
    r["frame_png_b64"] = images
    e = time.monotonic()
    r["export"] = s.json("command_run", id="file.exportMedia",
                         params={"path": f"{OUT}/sintel-720.mp4", "format": "h264", "width": 1280, "height": 544,
                                 "wait": True})
    r["export_seconds"] = round(time.monotonic() - e, 2)
    r["craft_seconds"] = round(time.monotonic() - t, 2)
    r["peak_kb"] = s.close()
    r["calls"] = s.log
    r["check"] = box(f"""ls -l {OUT}/sintel-720.mp4
python3 - <<'PY'
import struct
data=open('{OUT}/sintel-720.mp4','rb').read()
i=data.find(b'mvhd')
v=data[i+4]
ts,dur=(struct.unpack('>IQ',data[i+24:i+36]) if v==1 else struct.unpack('>II',data[i+16:i+24]))
print('mp4 duration seconds', round(dur/ts,2), '- avc1 track:', b'avc1' in data[:200000] or b'avc1' in data[-200000:])
PY
""")
    r["baseline"] = box("command -v ffmpeg ffprobe || echo 'no ffmpeg or ffprobe in the box: no server-side video path today'")
    return r


def photocraft():
    """A 4948x3786 Wikimedia photo of Ponte Vecchio: resize to 2400 px wide, lift brightness and
    contrast, export a JPEG, against Pillow doing the same."""
    s = Session("photocraft")
    r = {"file": "ponte-vecchio-dusk.jpg"}
    t = time.monotonic()
    r["open"] = s.json("doc_open", path=f"{REAL}/ponte-vecchio-dusk.jpg")
    r["size"] = s.json("command_run", id="image.imageSize", params={"width": 2400})
    r["adjust"] = s.json("command_run", id="image.adjustments.brightnessContrast",
                         params={"brightness": 25, "contrast": 10})
    s.call("doc_export", path=f"{OUT}/ponte-2400.jpg", quality=85)
    _, images, _ = s.call("render_preview", max_side=512)
    r["preview_png_b64"] = images
    r["craft_seconds"] = round(time.monotonic() - t, 2)
    r["peak_kb"] = s.close()
    r["calls"] = s.log
    r["baseline"] = box(f"""python3 - <<'PY'
import time
from PIL import Image, ImageEnhance, ImageStat
s=time.monotonic()
im=Image.open('{REAL}/ponte-vecchio-dusk.jpg'); im=im.resize((2400, round(im.height*2400/im.width)), Image.LANCZOS)
im=ImageEnhance.Contrast(ImageEnhance.Brightness(im).enhance(1.15)).enhance(1.08); im.save('/tmp/ponte-pillow.jpg', quality=85)
print('Pillow', round(time.monotonic()-s,2), 's')
for n in ['{REAL}/ponte-vecchio-dusk.jpg', '{OUT}/ponte-2400.jpg', '/tmp/ponte-pillow.jpg']:
    x=Image.open(n)
    print(n.rsplit('/',1)[1], x.size, x.mode, 'mean luma', round(ImageStat.Stat(x.convert('L')).mean[0],1), 'exif', bool(x.info.get('exif')))
PY
""")
    return r
