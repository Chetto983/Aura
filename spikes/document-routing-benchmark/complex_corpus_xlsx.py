"""Workbooks of the complex retrieval corpus: several sheets, merged headers, formulas.

openpyxl writes formulas without cached values; complex_corpus.py has LibreOffice recalculate
and re-save each workbook, as Excel would, so a reader sees both the formula and its value.
"""

from __future__ import annotations

import datetime
import random

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

FIXED_TIME = datetime.datetime(2026, 10, 16, 9, 0, 0)
REGIONS = ("Nord-Ovest", "Nord-Est", "Centro", "Sud e Isole")
QUARTERS = ("T1", "T2", "T3", "T4")
CENTRES = ("Produzione", "Logistica", "Commerciale", "Amministrazione", "R&S")


def _workbook() -> Workbook:
    wb = Workbook()
    wb.properties.creator = "Ufficio Controllo di Gestione"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    return wb


def _header(ws, row: int, labels: list[str], start_col: int = 1) -> None:
    for offset, label in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + offset, value=label)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def budget(year: int, seed: int) -> tuple[Workbook, dict]:
    """A consolidated budget: assumptions, revenue by region and quarter, costs, EBITDA."""
    rng = random.Random(seed)
    wb = _workbook()
    assumptions = wb.active
    assumptions.title = "Ipotesi"
    growth = round(rng.uniform(0.03, 0.07), 3)
    inflation = round(rng.uniform(0.015, 0.03), 3)
    assumptions.append([f"Ipotesi di budget {year}", None])
    assumptions.append(["Crescita ricavi attesa", growth])
    assumptions.append(["Inflazione sui costi", inflation])
    assumptions.append(["Cambio EUR/USD di riferimento", round(rng.uniform(1.05, 1.15), 2)])
    assumptions["A2"].comment = Comment("Approvata dal CdA del 12 dicembre", "CFO")

    revenue = wb.create_sheet("Ricavi")
    revenue.merge_cells("A1:F1")
    revenue["A1"] = f"Ricavi {year} per area geografica (migliaia di euro)"
    revenue.merge_cells("B2:E2")
    revenue["B2"] = "Trimestre"
    _header(revenue, 3, ["Area", *QUARTERS, "Totale anno"])
    totals: dict[str, int] = {}
    for index, region in enumerate(REGIONS, start=4):
        values = [rng.randrange(800, 2600) for _ in QUARTERS]
        revenue.append([region, *values, f"=SUM(B{index}:E{index})"])
        totals[region] = sum(values)
    last = 3 + len(REGIONS)
    revenue.append(["Totale", *[f"=SUM({c}4:{c}{last})" for c in "BCDEF"]])
    revenue.freeze_panes = "B4"

    costs = wb.create_sheet("Costi")
    _header(costs, 1, ["Centro di costo", "Personale", "Materiali", "Servizi", "Ammortamenti", "Totale"])
    cost_totals: dict[str, int] = {}
    for index, centre in enumerate(CENTRES, start=2):
        values = [rng.randrange(150, 1400) for _ in range(4)]
        costs.append([centre, *values, f"=SUM(B{index}:E{index})"])
        cost_totals[centre] = sum(values)
    end = 1 + len(CENTRES)
    costs.append(["Totale", *[f"=SUM({c}2:{c}{end})" for c in "BCDEF"]])

    summary = wb.create_sheet("Consolidato")
    summary.append(["Voce", "Importo (k€)"])
    summary.append(["Ricavi totali", f"=Ricavi!F{last + 1}"])
    summary.append(["Costi totali", f"=Costi!F{end + 1}"])
    summary.append(["Costi indicizzati", "=B3*(1+Ipotesi!B3)"])
    summary.append(["EBITDA", "=B2-B4"])
    summary.append(["EBITDA margin", "=B5/B2"])
    revenue_total = sum(totals.values())
    cost_total = sum(cost_totals.values())
    ebitda = revenue_total - cost_total * (1 + inflation)
    facts = {
        "region_totals": totals,
        "revenue_total": revenue_total,
        "cost_totals": cost_totals,
        "growth": growth,
        "inflation": inflation,
        "ebitda": round(ebitda),
        "top_region": max(totals, key=totals.get),
        "top_centre": max(cost_totals, key=cost_totals.get),
    }
    return wb, facts


def spare_parts(seed: int, rows: int = 1500) -> tuple[Workbook, dict]:
    """A spare-parts store: stock, suppliers looked up by formula, and a reorder list."""
    rng = random.Random(seed)
    wb = _workbook()
    suppliers = [
        ("F01", "Meccanica Bassi S.r.l.", "Brescia"),
        ("F02", "Rulmeca Italia", "Bergamo"),
        ("F03", "Elettrotecnica Ferri", "Torino"),
        ("F04", "Pneumax Service", "Lurano"),
        ("F05", "Guarnizioni Alpi", "Cuneo"),
    ]
    families = [
        "Cuscinetto",
        "Cinghia dentata",
        "Sensore induttivo",
        "Elettrovalvola",
        "Guarnizione",
        "Motoriduttore",
        "Fotocellula",
        "Encoder",
        "Lama di taglio",
        "Rullo trasportatore",
    ]
    stock = wb.active
    stock.title = "Giacenze"
    _header(
        stock,
        1,
        [
            "Codice",
            "Descrizione",
            "Ubicazione",
            "Fornitore",
            "Ragione sociale",
            "Giacenza",
            "Scorta minima",
            "Da riordinare",
        ],
    )
    validation = DataValidation(type="list", formula1='"F01,F02,F03,F04,F05"', allow_blank=False)
    stock.add_data_validation(validation)
    below: list[str] = []
    chosen = {}
    for index in range(rows):
        row = index + 2
        code = f"RIC-{10000 + index * 7}"
        family = families[index % len(families)]
        description = (
            f"{family} {rng.choice(['tipo A', 'tipo B', 'rinforzato', 'inox'])} {rng.randrange(10, 99)} mm"
        )
        location = f"S{rng.randrange(1, 9)}-{rng.choice('ABCDEF')}{rng.randrange(1, 30):02d}"
        supplier = suppliers[index % len(suppliers)][0]
        quantity = rng.randrange(0, 60)
        minimum = rng.randrange(5, 25)
        stock.append(
            [
                code,
                description,
                location,
                supplier,
                f"=VLOOKUP(D{row},Fornitori!$A$2:$C$6,2,FALSE)",
                quantity,
                minimum,
                f'=IF(F{row}<G{row},"SI","NO")',
            ]
        )
        validation.add(f"D{row}")
        if quantity < minimum:
            below.append(code)
        if index in (137, 911):
            chosen[code] = {
                "description": description,
                "location": location,
                "supplier": dict((s[0], s[1]) for s in suppliers)[supplier],
                "quantity": quantity,
            }
    stock.freeze_panes = "A2"
    supplier_sheet = wb.create_sheet("Fornitori")
    _header(supplier_sheet, 1, ["Codice", "Ragione sociale", "Città"])
    for supplier in suppliers:
        supplier_sheet.append(list(supplier))
    reorder = wb.create_sheet("Riordino")
    reorder["A1"] = "Articoli sotto scorta minima"
    reorder["A1"].font = Font(bold=True)
    reorder["E1"] = "Lead time medio per fornitore (giorni)"
    reorder["E1"].font = Font(bold=True)
    for offset, code in enumerate(below[:40], start=2):
        reorder.cell(row=offset, column=1, value=code)
    lead_times = {s[0]: rng.randrange(5, 40) for s in suppliers}
    for offset, (supplier, days) in enumerate(lead_times.items(), start=2):
        reorder.cell(row=offset, column=5, value=supplier)
        reorder.cell(row=offset, column=6, value=days)
    slowest = max(lead_times, key=lead_times.get)
    facts = {
        "rows": rows,
        "below_minimum": len(below),
        "chosen": chosen,
        "lead_times": lead_times,
        "slowest_supplier": dict((s[0], s[1]) for s in suppliers)[slowest],
    }
    return wb, facts


def maintenance_plan(seed: int) -> tuple[Workbook, dict]:
    """Machines against weeks, with a legend of activity codes on its own sheet."""
    rng = random.Random(seed)
    wb = _workbook()
    plan = wb.active
    plan.title = "Piano 2026"
    machines = [
        "Confezionatrice FX300",
        "Confezionatrice FX200",
        "Etichettatrice EL40",
        "Pallettizzatore PZ8",
        "Avvolgitore AV2",
        "Nastro di linea L3",
        "Compressore C55",
    ]
    plan.merge_cells("A1:A2")
    plan["A1"] = "Macchina"
    for month in range(12):
        first = 2 + month * 4
        plan.merge_cells(start_row=1, start_column=first, end_row=1, end_column=first + 3)
        plan.cell(row=1, column=first, value=datetime.date(2026, month + 1, 1).strftime("%m/%Y"))
        for week in range(4):
            plan.cell(row=2, column=first + week, value=f"S{month * 4 + week + 1}")
    codes = ["L", "I", "S", "R", "T"]
    revisions: dict[str, int] = {}
    for row, machine in enumerate(machines, start=3):
        plan.cell(row=row, column=1, value=machine)
        for column in range(2, 50):
            if rng.random() < 0.18:
                code = rng.choice(codes)
                plan.cell(row=row, column=column, value=code)
                if code == "R":
                    revisions[machine] = revisions.get(machine, 0) + 1
    legend = wb.create_sheet("Legenda")
    for entry in (
        ("L", "Lubrificazione"),
        ("I", "Ispezione visiva"),
        ("S", "Sostituzione parti usurate"),
        ("R", "Revisione generale con fermo linea"),
        ("T", "Taratura sensori"),
    ):
        legend.append(list(entry))
    most = max(revisions, key=revisions.get)
    return wb, {"revisions": revisions, "most_revisions": most}


def shifts(seed: int) -> tuple[Workbook, dict]:
    """A department's October rota: people against days, hours summed by formula."""
    rng = random.Random(seed)
    wb = _workbook()
    rota = wb.active
    rota.title = "Ottobre 2026"
    people = [
        "Rossi Marco",
        "Bianchi Laura",
        "Verdi Paolo",
        "Gallo Sara",
        "Conti Luca",
        "Costa Elena",
        "Fontana Davide",
        "Marino Giulia",
        "Greco Andrea",
        "Bruno Chiara",
    ]
    _header(rota, 1, ["Operatore", *[str(day) for day in range(1, 32)], "Ore totali"])
    nights: dict[str, int] = {}
    for row, person in enumerate(people, start=2):
        days = [rng.choice("MMPPNFR") for _ in range(31)]
        rota.append(
            [
                person,
                *days,
                f'=8*(COUNTIF(B{row}:AF{row},"M")+COUNTIF(B{row}:AF{row},"P")+COUNTIF(B{row}:AF{row},"N"))',
            ]
        )
        nights[person] = days.count("N")
    rota.append([])
    rota.append(["Legenda: M mattino 6-14, P pomeriggio 14-22, N notte 22-6, F ferie, R riposo"])
    most = max(nights, key=nights.get)
    return wb, {"nights": nights, "most_nights": most}
